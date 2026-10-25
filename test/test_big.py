from docx import Document
from pathlib import Path
import json
import re
import base64
import os


# ============================================================
# CONFIGURATION
# ============================================================

# Detects:
# _____
# __________
# ..........
# -----------
# ________ etc.
BLANK_PATTERN = re.compile(
    r"(_{3,}|\.{5,}|-{5,}|_{2,}\s*_{2,})"
)


# ============================================================
# BLANK DETECTION
# ============================================================

def detect_blanks(text):
    """
    Detect fillable blanks inside text.

    Returns a list containing the text segments and blanks.
    """

    if not text:
        return []

    result = []
    last_end = 0

    for match in BLANK_PATTERN.finditer(text):

        # Normal text before blank
        if match.start() > last_end:
            result.append({
                "type": "text",
                "text": text[last_end:match.start()]
            })

        # Blank
        result.append({
            "type": "blank",
            "placeholder": match.group(),
            "required": True
        })

        last_end = match.end()

    # Remaining text
    if last_end < len(text):
        result.append({
            "type": "text",
            "text": text[last_end:]
        })

    return result


# ============================================================
# PARAGRAPH TYPE
# ============================================================

def get_paragraph_type(paragraph):

    style = paragraph.style.name.lower() if paragraph.style else ""

    # Heading
    if "heading" in style:
        return "heading"

    # Lists
    if "list bullet" in style:
        return "bullet_list"

    if "list number" in style:
        return "numbered_list"

    # Normal paragraph
    return "paragraph"


# ============================================================
# EXTRACT PARAGRAPH
# ============================================================

def extract_paragraph(paragraph, paragraph_id):

    text = paragraph.text.strip()

    paragraph_type = get_paragraph_type(paragraph)

    result = {
        "id": paragraph_id,
        "type": paragraph_type,
        "style": paragraph.style.name if paragraph.style else None,
        "content": []
    }

    # Completely empty paragraph
    if not text:

        result["content"].append({
            "type": "empty_paragraph"
        })

        return result

    # Detect blanks
    content = detect_blanks(text)

    result["content"] = content

    return result


# ============================================================
# EXTRACT IMAGE
# ============================================================

def extract_images(doc, output_dir):

    image_dir = os.path.join(output_dir, "images")

    os.makedirs(image_dir, exist_ok=True)

    images = []

    for relationship in doc.part.rels.values():

        if "image" not in relationship.reltype:
            continue

        image_part = relationship.target_part

        image_name = os.path.basename(
            image_part.partname
        )

        image_path = os.path.join(
            image_dir,
            image_name
        )

        with open(image_path, "wb") as f:
            f.write(image_part.blob)

        images.append({
            "type": "image",
            "filename": image_name,
            "path": image_path
        })

    return images


# ============================================================
# TABLE CELL EXTRACTION
# ============================================================

def extract_cell(cell, row_index, column_index):

    text = cell.text.strip()

    # --------------------------------------------------------
    # EMPTY CELL
    # --------------------------------------------------------

    if not text:

        return {
            "row": row_index,
            "column": column_index,
            "type": "empty_cell",
            "fillable": True,
            "content": []
        }

    # --------------------------------------------------------
    # NORMAL / BLANK-CONTAINING CELL
    # --------------------------------------------------------

    content = detect_blanks(text)

    contains_blank = any(
        item["type"] == "blank"
        for item in content
    )

    return {
        "row": row_index,
        "column": column_index,
        "type": "fillable_cell" if contains_blank else "cell",
        "fillable": contains_blank,
        "content": content
    }


# ============================================================
# TABLE EXTRACTION
# ============================================================

def extract_table(table, table_id):

    result = {
        "id": table_id,
        "type": "table",
        "rows": []
    }

    for row_index, row in enumerate(table.rows):

        row_data = {
            "row": row_index,
            "cells": []
        }

        for column_index, cell in enumerate(row.cells):

            cell_data = extract_cell(
                cell,
                row_index,
                column_index
            )

            row_data["cells"].append(cell_data)

        result["rows"].append(row_data)

    return result


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_docx(input_file, output_dir):

    print("Opening:", input_file)

    doc = Document(input_file)

    os.makedirs(output_dir, exist_ok=True)

    result = {

        "document": {
            "filename": Path(input_file).name,
            "format": "DOCX"
        },

        "elements": [],

        "images": [],

        "statistics": {
            "paragraphs": 0,
            "headings": 0,
            "lists": 0,
            "tables": 0,
            "empty_cells": 0,
            "blanks": 0,
            "images": 0
        }
    }

    # ========================================================
    # PARAGRAPHS
    # ========================================================

    paragraph_id = 1

    for paragraph in doc.paragraphs:

        extracted = extract_paragraph(
            paragraph,
            paragraph_id
        )

        result["elements"].append(extracted)

        result["statistics"]["paragraphs"] += 1

        if extracted["type"] == "heading":
            result["statistics"]["headings"] += 1

        if extracted["type"] in [
            "bullet_list",
            "numbered_list"
        ]:
            result["statistics"]["lists"] += 1

        # Count blanks
        for item in extracted["content"]:

            if item["type"] == "blank":
                result["statistics"]["blanks"] += 1

        paragraph_id += 1

    # ========================================================
    # TABLES
    # ========================================================

    table_id = 1

    for table in doc.tables:

        extracted_table = extract_table(
            table,
            table_id
        )

        result["elements"].append(
            extracted_table
        )

        result["statistics"]["tables"] += 1

        # Count empty cells
        for row in extracted_table["rows"]:

            for cell in row["cells"]:

                if cell["type"] == "empty_cell":

                    result["statistics"]["empty_cells"] += 1

                if cell.get("fillable"):

                    for item in cell["content"]:

                        if item["type"] == "blank":
                            result["statistics"]["blanks"] += 1

        table_id += 1

    # ========================================================
    # IMAGES
    # ========================================================

    images = extract_images(
        doc,
        output_dir
    )

    result["images"] = images

    result["statistics"]["images"] = len(images)

    # ========================================================
    # SAVE JSON
    # ========================================================

    json_path = os.path.join(
        output_dir,
        "extracted_document.json"
    )

    with open(
        json_path,
        "w",
        encoding="utf-8"
    ) as f:

        json.dump(
            result,
            f,
            indent=4,
            ensure_ascii=False
        )

    print("\nExtraction complete!")

    print(
        "JSON saved to:",
        json_path
    )

    print(
        "Images saved to:",
        os.path.join(
            output_dir,
            "images"
        )
    )

    print("\nStatistics:")

    for key, value in result["statistics"].items():

        print(
            f"{key}: {value}"
        )


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    input_file = "Exposure.docx"

    output_directory = "extracted_output"

    extract_docx(
        input_file,
        output_directory
    )