"""
docx_extractor.py
==================
Extracts fillable structure from a .docx form template.

Design goals:
  - Every blank/empty field gets a STABLE, DETERMINISTIC locator that
    docx_writer.py can use later to surgically insert a value without
    touching surrounding formatting (see docx_writer.py for the write-back
    half).
  - Labels/section/field_type are inferred locally (no network calls) so
    an admin only has to *correct* a few fields in the dashboard instead
    of typing every field from scratch.
  - Paragraphs inside table cells and paragraphs in the document body are
    processed by the exact same function, so locators for both look and
    behave the same way.

Output: extracted_document.json with two views of the same data —
  "elements": full structural tree (for rendering a document preview)
  "fields":   flat list of everything fillable (for the admin dashboard)
"""

from docx import Document
from pathlib import Path
import json
import re
import os


# ============================================================
# CONFIGURATION
# ============================================================

# Two blank conventions found in real templates:
#   underline-style:  Name: __________          (write directly over it)
#   bracket-style:    Company Name: [Your Company Name, Inc.]  (the bracket
#                      text itself is a hint/example, common in "sample" docs)
# Matched together, in document order, via named groups.
BLANK_PATTERN = re.compile(
    r"(?P<underline>_{3,}|\.{5,}|-{5,}|_{2,}\s*_{2,})"
    r"|\[(?P<bracket>[^\[\]\n]{1,60})\]"
)

# Checkbox / option glyphs used for single- or multi-select fields in Word
# forms (Wingdings checkboxes, plain box-drawing, or checkmarks).
CHECKBOX_GLYPHS = "☐☑☒□■✓✔"

# Local, offline field-type classifier, checked against a resolved LABEL.
# Order matters — first match wins. Extended to roughly cover the common
# form-field taxonomy (text/numeric/choice/date/contact/upload/rating).
FIELD_TYPE_RULES = [
    (re.compile(r"\bdate\b|\bdob\b|\bd\.o\.b\b|\bmm[/\-.]dd[/\-.]yyyy\b", re.I), "date"),
    (re.compile(r"\btime\b(?!stamp)", re.I), "time"),
    (re.compile(r"\bsign(ature)?\b|\binitial(s)?\b", re.I), "signature"),
    (re.compile(r"\be-?mail\b", re.I), "email"),
    (re.compile(r"\b(phone|contact\s*no|mobile|tel\.?)\b", re.I), "phone"),
    (re.compile(r"\burl\b|\bwebsite\b|\bhyperlink\b", re.I), "url"),
    (re.compile(r"\baddress\b", re.I), "address"),
    (re.compile(r"\b(percent|percentage|%)\b", re.I), "percentage"),
    (re.compile(r"\b(amount|salary|price|cost|fee|currency|budget)\b", re.I), "currency"),
    (re.compile(r"\b(qty|quantity|no\.?|number|age|count)\b", re.I), "number"),
    (re.compile(r"\b(upload|attach(ment)?|file)\b", re.I), "file_upload"),
    (re.compile(r"\b(photo|image|picture)\b", re.I), "image_upload"),
    (re.compile(r"\brating\b|\bstars?\b", re.I), "rating"),
    (re.compile(
        r"\b(summary|description|remarks?|notes?|comments?|findings|details|"
        r"explanation|justification|discussion|recommendation|commitment)\b", re.I
    ), "multiline_text"),
]

# Placeholder-hint rules — applied to bracket-style content, e.g. the "MM/DD/YYYY"
# inside "[MM/DD/YYYY]", which often reveals the type even with no label at all.
PLACEHOLDER_HINT_RULES = [
    (re.compile(r"^\s*mm[/\-.]dd[/\-.]yyyy\s*$", re.I), "date"),
    (re.compile(r"^\s*dd[/\-.]mm[/\-.]yyyy\s*$", re.I), "date"),
    (re.compile(r"^\s*\d+\s*$"), "number"),
    (re.compile(r"e-?mail", re.I), "email"),
    (re.compile(r"address", re.I), "address"),
    (re.compile(r"date", re.I), "date"),
    (re.compile(r"name,?\s*title", re.I), "text"),
]

LABEL_STRIP_CHARS = ":-–—.\u2022 \t"

# A short, mostly-uppercase standalone line (e.g. "PROPONENT'S NAME") reads as
# a caption, not a sentence — used to decide whether it's safe to use a
# paragraph BELOW a lone blank as that blank's label.
def looks_like_caption(text):
    text = text.strip()
    if not text or len(text) > 40:
        return False
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return False
    upper_ratio = sum(c.isupper() for c in letters) / len(letters)
    return upper_ratio > 0.6


# ============================================================
# BLANK DETECTION (with character offsets — needed for write-back)
# ============================================================

def detect_blanks(text):
    """
    Split `text` into a list of {"type": "text"/"blank", ...} segments.
    Blank segments carry char_start/char_end offsets into the ORIGINAL
    `text` string so docx_writer.py can map them back onto real docx runs.
    """
    if not text:
        return []

    result = []
    last_end = 0

    for match in BLANK_PATTERN.finditer(text):
        if match.start() > last_end:
            result.append({
                "type": "text",
                "text": text[last_end:match.start()],
                "char_start": last_end,
                "char_end": match.start(),
            })

        bracket_hint = match.group("bracket")
        result.append({
            "type": "blank",
            "style": "bracket" if bracket_hint is not None else "underline",
            "placeholder": match.group(),
            "hint": bracket_hint,  # e.g. "MM/DD/YYYY" — None for underline-style
            "char_start": match.start(),
            "char_end": match.end(),
        })

        last_end = match.end()

    if last_end < len(text):
        result.append({
            "type": "text",
            "text": text[last_end:],
            "char_start": last_end,
            "char_end": len(text),
        })

    return result


# ============================================================
# LABEL + FIELD-TYPE INFERENCE (local heuristics, no network)
# ============================================================

def clean_label(raw):
    if not raw:
        return None
    label = raw.strip().strip(LABEL_STRIP_CHARS).strip()
    return label or None


def infer_field_type(label, hint=None):
    """
    Try the placeholder hint first — "[MM/DD/YYYY]" is a stronger, more
    literal type signal than any label ever is — then fall back to the
    label-based rules, then plain text.
    """
    if hint:
        for pattern, field_type in PLACEHOLDER_HINT_RULES:
            if pattern.search(hint):
                return field_type

    if label:
        for pattern, field_type in FIELD_TYPE_RULES:
            if pattern.search(label):
                return field_type

    return "text"


def label_from_content(content, blank_index):
    """
    Rule 1: same-line label — nearest non-empty 'text' segment BEFORE
    this blank, within the same paragraph/cell content list.
    """
    for item in reversed(content[:blank_index]):
        if item["type"] == "text" and item["text"].strip():
            return clean_label(item["text"])
    return None


def label_from_previous_paragraph(elements, current_index):
    """
    Rule 3: line-above fallback — used when a paragraph is JUST a blank
    with no inline label text of its own (e.g. lone signature line).
    """
    for prev in reversed(elements[:current_index]):
        if prev.get("element_type") != "paragraph":
            continue
        text_parts = [
            seg["text"] for seg in prev["content"]
            if seg["type"] == "text" and seg["text"].strip()
        ]
        if text_parts:
            return clean_label(" ".join(text_parts))
        if any(seg["type"] == "blank" for seg in prev["content"]):
            # previous line is also just a blank — keep looking further up
            continue
        break
    return None


def label_from_next_paragraph(elements, current_index):
    """
    Rule: caption-below fallback — some forms put the blank FIRST and a
    short caption on the line right after it (e.g. a signature-style rule
    with "PROPONENT'S NAME" printed underneath). Only fires when the very
    next paragraph looks like a caption (short, mostly-caps) rather than
    ordinary prose, so it doesn't accidentally grab a following sentence
    or heading.
    """
    if current_index + 1 >= len(elements):
        return None
    nxt = elements[current_index + 1]
    if nxt.get("element_type") != "paragraph":
        return None
    text_parts = [
        seg["text"] for seg in nxt["content"]
        if seg["type"] == "text" and seg["text"].strip()
    ]
    if not text_parts:
        return None
    candidate = " ".join(text_parts)
    if looks_like_caption(candidate):
        return clean_label(candidate)
    return None


def label_from_table_headers(table_rows, row_index, column_index):
    """
    Rule 2: table context — try the column-0 cell of this row (row label),
    then the row-0 cell of this column (column header). Picks whichever
    has more non-empty text, since layout convention varies by document.
    """
    candidates = []

    if column_index != 0 and row_index < len(table_rows):
        row_label_cell = table_rows[row_index]["cells"][0]
        text = _cell_plaintext(row_label_cell)
        if text:
            candidates.append(text)

    if row_index != 0 and len(table_rows) > 0:
        header_cell = table_rows[0]["cells"][column_index] if column_index < len(table_rows[0]["cells"]) else None
        if header_cell:
            text = _cell_plaintext(header_cell)
            if text:
                candidates.append(text)

    if not candidates:
        return None

    return clean_label(max(candidates, key=len))


def _cell_plaintext(cell_dict):
    parts = []
    for para in cell_dict.get("paragraphs", []):
        for seg in para["content"]:
            if seg["type"] == "text" and seg["text"].strip():
                parts.append(seg["text"].strip())
    return " ".join(parts)


# ============================================================
# PARAGRAPH PROCESSING (shared by body paragraphs AND cell paragraphs)
# ============================================================

def get_paragraph_type(paragraph):
    style = paragraph.style.name.lower() if paragraph.style else ""
    if "heading" in style:
        return "heading"
    if "list bullet" in style:
        return "bullet_list"
    if "list number" in style:
        return "numbered_list"
    return "paragraph"


def process_paragraph(paragraph, locator):
    """
    Extract one paragraph (body OR table-cell) into a structured dict.
    `locator` is the caller-supplied dict identifying WHERE this paragraph
    lives in the document — it gets embedded verbatim into every blank
    so docx_writer.py knows exactly what to re-open and edit.

    IMPORTANT: offsets are computed against the RAW paragraph.text, not a
    stripped version. docx_writer.py's run-splitting logic operates against
    the unstripped run concatenation (which IS paragraph.text) — stripping
    here while keeping offsets would silently misalign every replacement by
    the amount of leading whitespace.
    """
    raw_text = paragraph.text
    paragraph_type = get_paragraph_type(paragraph)

    result = {
        "element_type": "paragraph",
        "type": paragraph_type,
        "style": paragraph.style.name if paragraph.style else None,
        "locator": locator,
        "content": [],
    }

    if not raw_text.strip():
        result["content"] = [{"type": "empty_paragraph"}]
        return result

    # Leading checkbox glyph -> this paragraph is a single selectable option
    # ("☐ During ECC Application Stage"). Split it off as its own segment,
    # then run normal blank detection on whatever text follows (an option
    # can still contain its own inline blank, e.g. "...approved on ____").
    stripped_leading = raw_text.lstrip()
    if stripped_leading and stripped_leading[0] in CHECKBOX_GLYPHS:
        glyph_offset = len(raw_text) - len(stripped_leading)
        glyph = stripped_leading[0]
        rest_start = glyph_offset + 1
        result["content"] = [{
            "type": "checkbox_option",
            "glyph": glyph,
            "checked": glyph in "☑☒✓✔",
            "char_start": glyph_offset,
            "char_end": rest_start,
        }]
        rest_segments = detect_blanks(raw_text[rest_start:])
        for seg in rest_segments:
            seg["char_start"] += rest_start
            seg["char_end"] += rest_start
        result["content"].extend(rest_segments)
        return result

    result["content"] = detect_blanks(raw_text)
    return result


def finalize_paragraph_fields(paragraph_dict, elements, element_index, section, fields_out):
    """
    Second pass: now that `content` exists, resolve label / field_type /
    section for every blank (and checkbox option) in this paragraph and
    append flat field record(s) to `fields_out`.
    """
    content = paragraph_dict["content"]
    is_pure_blank_line = (
        len(content) == 1
        and content[0]["type"] == "blank"
        and content[0]["char_start"] == 0
    )
    blank_n = 0

    for i, item in enumerate(content):

        if item["type"] == "checkbox_option":
            # Label = whatever text segments follow the glyph on this line.
            option_text = " ".join(
                seg["text"] for seg in content[i + 1:]
                if seg["type"] == "text" and seg["text"].strip()
            ).strip()
            label = clean_label(option_text) or None

            field_id = f'{paragraph_dict["locator"]["ref"]}-cb'
            fields_out.append({
                "field_id": field_id,
                "kind": "checkbox_option",
                "locator": {**paragraph_dict["locator"], "char_start": item["char_start"], "char_end": item["char_end"]},
                "label": label,
                "field_type": "checkbox",
                "section": section,
                "placeholder": item["glyph"],
                "required": False,
                "checked_default": item["checked"],
                # group_id is filled in by assign_choice_groups() afterwards
                "group_id": None,
            })
            item["field_id"] = field_id
            continue

        if item["type"] != "blank":
            continue

        label = label_from_content(content, i)
        if label is None and is_pure_blank_line:
            # Try "caption below" before "text above" — a lone blank line
            # is exactly the case the caption-below convention shows up in.
            label = label_from_next_paragraph(elements, element_index)
        if label is None:
            label = label_from_previous_paragraph(elements, element_index)

        field_id = f'{paragraph_dict["locator"]["ref"]}-b{blank_n}'

        fields_out.append({
            "field_id": field_id,
            "kind": "paragraph_blank",
            "locator": {**paragraph_dict["locator"], "char_start": item["char_start"], "char_end": item["char_end"]},
            "label": label,
            "field_type": infer_field_type(label, item.get("hint")),
            "section": section,
            "placeholder": item["placeholder"],
            "hint": item.get("hint"),
            "required": True,
        })

        item["field_id"] = field_id
        blank_n += 1


# ============================================================
# POST-PROCESSING PASSES (operate on body-level paragraphs)
# ============================================================

def merge_multiline_blank_groups(elements, fields):
    """
    Several documents represent a long-answer box as 2+ consecutive
    paragraphs that are EACH just a full-line blank
    ("Changes in Project Design (if any):" followed by three blank
    underscore lines). Left alone, that produces three duplicate fields
    all fighting for the same label. This merges any such run into one
    "multiline_blank_group" field with a locator per line, so a single
    answer can be written across all of them (see docx_writer.py).
    """
    body_paragraphs = [
        (i, e) for i, e in enumerate(elements)
        if e["element_type"] == "paragraph" and e["locator"]["container"] == "body"
    ]

    def is_pure_blank(el):
        c = el["content"]
        return len(c) == 1 and c[0]["type"] == "blank" and c[0]["char_start"] == 0

    i = 0
    while i < len(body_paragraphs):
        idx, el = body_paragraphs[i]
        if not is_pure_blank(el):
            i += 1
            continue

        run = [(idx, el)]
        j = i + 1
        while j < len(body_paragraphs) and is_pure_blank(body_paragraphs[j][1]):
            run.append(body_paragraphs[j])
            j += 1

        if len(run) >= 2:
            first_idx = run[0][0]
            label = label_from_previous_paragraph(elements, first_idx)

            member_field_ids = {f'p{ridx}-b0' for ridx, _ in run}
            fields[:] = [f for f in fields if f["field_id"] not in member_field_ids]

            locators = []
            for ridx, rel in run:
                blank_seg = rel["content"][0]
                locators.append({
                    **rel["locator"],
                    "char_start": blank_seg["char_start"],
                    "char_end": blank_seg["char_end"],
                })

            merged_field_id = f'p{run[0][0]}__p{run[-1][0]}-multiline'
            fields.append({
                "field_id": merged_field_id,
                "kind": "multiline_blank_group",
                "locators": locators,
                "label": label,
                "field_type": "multiline_text",
                "section": None,
                "placeholder": None,
                "required": True,
                "line_count": len(run),
            })

        i = j


def assign_choice_groups(elements, fields):
    """
    Tags consecutive checkbox_option fields (no non-checkbox paragraph
    between them) with a shared group_id, so the dashboard can render
    them as one choice group instead of N unrelated checkboxes. Does NOT
    change how each option is written back — every option stays its own
    independently-toggleable field_id.
    """
    body_paragraphs = [
        e for e in elements
        if e["element_type"] == "paragraph" and e["locator"]["container"] == "body"
    ]
    fields_by_id = {f["field_id"]: f for f in fields}

    group_n = 0
    current_group = []

    def flush():
        nonlocal group_n, current_group
        if len(current_group) >= 1:
            gid = f"cg{group_n}"
            for fid in current_group:
                fields_by_id[fid]["group_id"] = gid
            group_n += 1
        current_group = []

    for el in body_paragraphs:
        checkbox_field_id = next(
            (seg.get("field_id") for seg in el["content"] if seg["type"] == "checkbox_option"),
            None,
        )
        if checkbox_field_id:
            current_group.append(checkbox_field_id)
        else:
            flush()
    flush()


# ============================================================
# TABLE PROCESSING
# ============================================================

def process_table(table, table_index):
    """
    Pass 1: build the raw row/cell/paragraph structure (no labels yet —
    label inference for cells needs the WHOLE table, so it happens after
    every cell's paragraphs are extracted).
    """
    result = {
        "element_type": "table",
        "locator": {"ref": f"t{table_index}", "container": "table", "table_index": table_index},
        "rows": [],
    }

    for row_index, row in enumerate(table.rows):
        row_data = {"row": row_index, "cells": []}

        for col_index, cell in enumerate(row.cells):
            cell_locator_base = {
                "container": "cell",
                "table_index": table_index,
                "row": row_index,
                "col": col_index,
            }

            cell_paragraphs = []
            for p_index, para in enumerate(cell.paragraphs):
                para_locator = {
                    **cell_locator_base,
                    "ref": f"t{table_index}-r{row_index}-c{col_index}-p{p_index}",
                    "para_index": p_index,
                }
                cell_paragraphs.append(process_paragraph(para, para_locator))

            has_text = bool(cell.text.strip())
            row_data["cells"].append({
                "row": row_index,
                "column": col_index,
                "type": "cell" if has_text else "empty_cell",
                "paragraphs": cell_paragraphs,
                # locator for the "whole empty cell" case (fill first paragraph)
                "empty_cell_locator": {
                    **cell_locator_base,
                    "ref": f"t{table_index}-r{row_index}-c{col_index}-p0",
                    "para_index": 0,
                },
            })

        result["rows"].append(row_data)

    return result


def finalize_table_fields(table_dict, fields_out):
    """
    Pass 2: now that all cells are extracted, resolve labels using
    row-0 / column-0 context, and emit flat field records for:
      - blanks found inside cell text  (fillable_cell)
      - cells that are entirely empty  (empty_cell)
    """
    rows = table_dict["rows"]

    for row in rows:
        for cell in row["cells"]:
            row_i, col_i = cell["row"], cell["column"]

            # Header row cells are context, not fields — skip row 0 by
            # convention (first row = column headers). We deliberately do
            # NOT blanket-skip column 0: it's only a "label column" in
            # some layouts, and in others (like this test case) it's a
            # normal data column whose row-1 cell should stay fillable.
            # If a document really uses a label column, those cells will
            # have their own text and get excluded by the empty_cell
            # check below anyway.
            is_structural_label_cell = (row_i == 0)

            any_blank_in_cell = False
            for p_index, para_dict in enumerate(cell["paragraphs"]):
                blank_n = 0
                content = para_dict["content"]
                for i, item in enumerate(content):
                    if item["type"] != "blank":
                        continue
                    any_blank_in_cell = True

                    label = label_from_content(content, i)
                    if label is None:
                        label = label_from_table_headers(rows, row_i, col_i)

                    field_id = f'{para_dict["locator"]["ref"]}-b{blank_n}'
                    fields_out.append({
                        "field_id": field_id,
                        "kind": "cell_blank",
                        "locator": {**para_dict["locator"], "char_start": item["char_start"], "char_end": item["char_end"]},
                        "label": label,
                        "field_type": infer_field_type(label, item.get("hint")),
                        "section": None,
                        "placeholder": item["placeholder"],
                        "hint": item.get("hint"),
                        "required": True,
                    })
                    item["field_id"] = field_id
                    blank_n += 1

            if cell["type"] == "empty_cell" and not is_structural_label_cell:
                label = label_from_table_headers(rows, row_i, col_i)
                field_id = cell["empty_cell_locator"]["ref"]
                fields_out.append({
                    "field_id": field_id,
                    "kind": "empty_cell",
                    "locator": cell["empty_cell_locator"],
                    "label": label,
                    # Wide/descriptive column headers ("Summary of Actions
                    # Taken", "Remarks") usually mean the cell wants a
                    # paragraph of writing, not a one-line answer.
                    "field_type": infer_field_type(label),
                    "section": None,
                    "placeholder": None,
                    "required": True,
                })


# ============================================================
# IMAGE EXTRACTION (filename-collision safe)
# ============================================================

def extract_images(doc, output_dir):
    image_dir = os.path.join(output_dir, "images")
    os.makedirs(image_dir, exist_ok=True)

    images = []
    for i, relationship in enumerate(doc.part.rels.values()):
        if "image" not in relationship.reltype:
            continue

        image_part = relationship.target_part
        original_name = os.path.basename(image_part.partname)
        safe_name = f"{i:03d}_{original_name}"
        image_path = os.path.join(image_dir, safe_name)

        with open(image_path, "wb") as f:
            f.write(image_part.blob)

        images.append({"type": "image", "filename": safe_name, "path": image_path})

    return images


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_docx(input_file, output_dir):
    print("Opening:", input_file)
    doc = Document(input_file)
    os.makedirs(output_dir, exist_ok=True)

    elements = []
    fields = []
    current_section = None

    # ---- Pass 1: body paragraphs — build structure + track section ----
    for para_index, paragraph in enumerate(doc.paragraphs):
        locator = {"ref": f"p{para_index}", "container": "body", "para_index": para_index}
        para_dict = process_paragraph(paragraph, locator)
        elements.append(para_dict)

        if para_dict["type"] == "heading":
            heading_text = paragraph.text.strip()
            if heading_text:
                current_section = heading_text

        para_dict["_section_at_this_point"] = current_section

    # ---- Pass 2: resolve labels for body paragraph blanks ----
    for element_index, para_dict in enumerate(elements):
        if para_dict["element_type"] != "paragraph":
            continue
        finalize_paragraph_fields(
            para_dict, elements, element_index,
            para_dict.pop("_section_at_this_point"), fields
        )

    # ---- Post-process body paragraphs: merge multiline blank runs, ----
    # ---- group consecutive checkboxes ----
    merge_multiline_blank_groups(elements, fields)
    assign_choice_groups(elements, fields)

    # ---- Tables ----
    table_dicts = []
    for table_index, table in enumerate(doc.tables):
        table_dict = process_table(table, table_index)
        finalize_table_fields(table_dict, fields)
        elements.append(table_dict)
        table_dicts.append(table_dict)

    # ---- Images ----
    images = extract_images(doc, output_dir)

    # ---- Statistics ----
    statistics = {
        "paragraphs": sum(1 for e in elements if e["element_type"] == "paragraph"),
        "headings": sum(1 for e in elements if e["element_type"] == "paragraph" and e["type"] == "heading"),
        "lists": sum(1 for e in elements if e["element_type"] == "paragraph" and e["type"] in ("bullet_list", "numbered_list")),
        "tables": len(table_dicts),
        "fields_total": len(fields),
        "fields_by_kind": {},
        "fields_by_type": {},
        "images": len(images),
    }
    for f in fields:
        statistics["fields_by_kind"][f["kind"]] = statistics["fields_by_kind"].get(f["kind"], 0) + 1
        statistics["fields_by_type"][f["field_type"]] = statistics["fields_by_type"].get(f["field_type"], 0) + 1

    result = {
        "document": {"filename": Path(input_file).name, "format": "DOCX"},
        "elements": elements,
        "fields": fields,
        "images": images,
        "statistics": statistics,
    }

    json_path = os.path.join(output_dir, "extracted_document.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, ensure_ascii=False)

    print("\nExtraction complete!")
    print("JSON saved to:", json_path)
    print("Images saved to:", os.path.join(output_dir, "images"))
    print("\nStatistics:")
    for key, value in statistics.items():
        print(f"{key}: {value}")

    return result


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    import sys
    input_file = sys.argv[1] if len(sys.argv) > 1 else "document.docx"
    output_directory = sys.argv[2] if len(sys.argv) > 2 else "extracted_output"
    extract_docx(input_file, output_directory)
