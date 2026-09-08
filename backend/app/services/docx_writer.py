"""
docx_writer.py
==============
Fills a .docx template using field values keyed by the field_id produced
by docx_extractor.py. This is the write-back half: it never touches
paragraph.text (read-only flattened view) — it edits the actual w:r runs
so surrounding formatting (bold, font, color) survives.

Typical flow:
    1. docx_extractor.py produces extracted_document.json (server stores it)
    2. admin/worker fills in {field_id: value} via the dashboard
    3. server calls fill_docx(original_path, output_path, values, fields)

`fields` is the flat "fields" list from extracted_document.json — each
field's "locator" says exactly which paragraph (body or table cell) and
which character span to replace.
"""

from docx import Document
from docx.text.paragraph import Paragraph


# ============================================================
# LOCATOR RESOLUTION
# ============================================================

def resolve_paragraph(doc, locator):
    """Return the python-docx Paragraph object a locator points to."""
    if locator["container"] == "body":
        return doc.paragraphs[locator["para_index"]]

    if locator["container"] == "cell":
        table = doc.tables[locator["table_index"]]
        cell = table.rows[locator["row"]].cells[locator["col"]]
        p_index = locator["para_index"]
        # Cell may not have that many paragraphs yet (e.g. was truly empty) —
        # pad with empty paragraphs so writing to para_index 0 always works.
        while len(cell.paragraphs) <= p_index:
            cell.add_paragraph("")
        return cell.paragraphs[p_index]

    raise ValueError(f"Unknown locator container: {locator['container']}")


# ============================================================
# RUN-LEVEL TEXT REPLACEMENT
# ============================================================

def replace_span(paragraph: Paragraph, char_start, char_end, new_text):
    """
    Replace paragraph.text[char_start:char_end] with new_text, editing
    the underlying runs in place so formatting of surrounding text is
    untouched. New text inherits the formatting of the first run it
    overlaps.

    If char_start == char_end (e.g. filling a genuinely empty paragraph,
    both 0), new_text is inserted there.
    """
    runs = paragraph.runs

    if not runs:
        # Paragraph had no runs at all (e.g. was completely empty).
        paragraph.add_run(new_text)
        return

    # Build cumulative offsets for each run.
    spans = []
    pos = 0
    for r in runs:
        length = len(r.text)
        spans.append((pos, pos + length))
        pos += length

    total_len = pos
    if char_start > total_len:
        # Locator stale relative to current doc state — append safely.
        runs[-1].text = runs[-1].text + new_text
        return

    first_idx = None
    last_idx = None
    for i, (s, e) in enumerate(spans):
        if e > char_start and first_idx is None:
            first_idx = i
        if s < char_end:
            last_idx = i

    if first_idx is None:
        # Blank sits exactly at end of paragraph text (char_start == total_len).
        runs[-1].text = runs[-1].text + new_text
        return
    if last_idx is None:
        last_idx = first_idx

    first_run = runs[first_idx]
    first_start = spans[first_idx][0]
    prefix = first_run.text[: char_start - first_start]

    if first_idx == last_idx:
        suffix = first_run.text[char_end - first_start:]
        first_run.text = prefix + new_text + suffix
    else:
        first_run.text = prefix + new_text

        for i in range(first_idx + 1, last_idx):
            runs[i].text = ""

        last_run = runs[last_idx]
        last_start = spans[last_idx][0]
        last_run.text = last_run.text[char_end - last_start:]


# ============================================================
# TOP-LEVEL FILL
# ============================================================

CHECKED_GLYPH = "☑"
UNCHECKED_GLYPH = "☐"


def fill_docx(input_path, output_path, values, fields):
    """
    values: dict of {field_id: value}  — from the dashboard/worker.
            Most fields take a string. checkbox_option fields take a
            truthy/falsy value (bool, "true"/"false", "yes"/"no", 1/0).
    fields: the "fields" list from extracted_document.json
             (needed to know each field_id's locator(s) + kind)
    """
    doc = Document(input_path)
    fields_by_id = {f["field_id"]: f for f in fields}

    skipped = []
    filled = []

    # Group SINGLE-locator paragraph/cell blanks by their target paragraph
    # so multiple blanks in the same paragraph are replaced in one pass,
    # RIGHT TO LEFT by char offset — replacing left-to-right would shift
    # the offsets of blanks later in the same paragraph.
    by_paragraph = {}
    empty_cell_fields = []
    checkbox_fields = []
    multiline_group_fields = []

    for field_id, value in values.items():
        field = fields_by_id.get(field_id)
        if field is None:
            skipped.append({"field_id": field_id, "reason": "unknown field_id"})
            continue

        kind = field["kind"]

        if kind == "empty_cell":
            empty_cell_fields.append((field, value))
        elif kind == "checkbox_option":
            checkbox_fields.append((field, value))
        elif kind == "multiline_blank_group":
            multiline_group_fields.append((field, value))
        else:
            locator = field["locator"]
            para_key = _paragraph_key(locator)
            by_paragraph.setdefault(para_key, []).append((locator, value, field_id))

    for para_key, edits in by_paragraph.items():
        locator_template = edits[0][0]
        paragraph = resolve_paragraph(doc, locator_template)
        # right-to-left so earlier offsets stay valid as we edit
        edits_sorted = sorted(edits, key=lambda e: e[0]["char_start"], reverse=True)
        for locator, value, field_id in edits_sorted:
            replace_span(paragraph, locator["char_start"], locator["char_end"], value)
            filled.append(field_id)

    for field, value in empty_cell_fields:
        paragraph = resolve_paragraph(doc, field["locator"])
        replace_span(paragraph, 0, len(paragraph.text), value)
        filled.append(field["field_id"])

    for field, value in checkbox_fields:
        paragraph = resolve_paragraph(doc, field["locator"])
        loc = field["locator"]
        glyph = CHECKED_GLYPH if _is_truthy(value) else UNCHECKED_GLYPH
        replace_span(paragraph, loc["char_start"], loc["char_end"], glyph)
        filled.append(field["field_id"])

    for field, value in multiline_group_fields:
        locators = field["locators"]
        # First line gets the value; every subsequent line in the group is
        # blanked out (its underscores removed) so the answer doesn't read
        # as "value" followed by leftover blank rules.
        first_para = resolve_paragraph(doc, locators[0])
        replace_span(first_para, locators[0]["char_start"], locators[0]["char_end"], value)
        for loc in locators[1:]:
            para = resolve_paragraph(doc, loc)
            replace_span(para, loc["char_start"], loc["char_end"], "")
        filled.append(field["field_id"])

    doc.save(output_path)
    return {"filled": filled, "skipped": skipped, "output": output_path}


def _is_truthy(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    if isinstance(value, str):
        return value.strip().lower() in ("true", "yes", "y", "1", "checked", "on")
    return bool(value)


def _paragraph_key(locator):
    if locator["container"] == "body":
        return ("body", locator["para_index"])
    return ("cell", locator["table_index"], locator["row"], locator["col"], locator["para_index"])


# ============================================================
# CLI (for local testing / handshake debugging)
# ============================================================

if __name__ == "__main__":
    import sys
    import json

    if len(sys.argv) < 5:
        print("Usage: python docx_writer.py <input.docx> <extracted.json> <values.json> <output.docx>")
        sys.exit(1)

    input_docx, extracted_json_path, values_json_path, output_docx = sys.argv[1:5]

    with open(extracted_json_path, encoding="utf-8") as f:
        extracted = json.load(f)
    with open(values_json_path, encoding="utf-8") as f:
        values = json.load(f)

    result = fill_docx(input_docx, output_docx, values, extracted["fields"])
    print(json.dumps(result, indent=2))
