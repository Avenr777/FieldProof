"""
generate_templates.py
Creates 4 realistic .docx document templates for FieldProof (Electrical, HVAC, Plumbing, Fire Safety)
using python-docx, with proper blanks, underlines, checkboxes, and tables.
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_electrical_template(filepath: str):
    doc = Document()

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("ELECTRICAL INSPECTION & COMPLIANCE CERTIFICATE")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(242, 98, 42)  # Brand orange

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("FieldProof Certified Electrical Documentation")
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph() # Spacer

    # Section 1: Customer & Site Details
    doc.add_heading("1. Site & Customer Details", level=2)
    doc.add_paragraph("Customer Name: _____________________")
    doc.add_paragraph("Site Address: _____________________")
    doc.add_paragraph("Inspection Date: [MM/DD/YYYY]")
    doc.add_paragraph("Technician Name: _____________________")

    doc.add_paragraph()

    # Section 2: Measurements Table
    doc.add_heading("2. Electrical System Measurements", level=2)
    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Parameter Description"
    hdr[1].text = "Expected Range"
    hdr[2].text = "Measured Value"

    rows_data = [
        ("Main Supply Voltage (V)", "220V - 240V", "_______"),
        ("Earth Resistance (Ohms)", "< 5.0 Ω", "_______"),
        ("Insulation Resistance (MΩ)", "> 1.0 MΩ", "_______"),
        ("Circuit Breaker Rating (Amps)", "10A - 63A", "_______"),
    ]
    for idx, (desc, exp, blank) in enumerate(rows_data, start=1):
        cells = table.rows[idx].cells
        cells[0].text = desc
        cells[1].text = exp
        cells[2].text = blank

    doc.add_paragraph()

    # Section 3: Safety Checklist
    doc.add_heading("3. Safety & Compliance Checklist", level=2)
    doc.add_paragraph("☐ Main Service Panel Grounded and Bonded")
    doc.add_paragraph("☐ Residual Current Device (RCD/GFCI) Operational")
    doc.add_paragraph("☐ No Damaged or Exposed Live Wiring")
    doc.add_paragraph("☐ Overcurrent Protection Devices Verified")

    doc.add_paragraph()

    # Section 4: Remarks & Sign-off
    doc.add_heading("4. Inspection Remarks & Authorization", level=2)
    doc.add_paragraph("Remarks: [Enter detailed inspection notes and recommendations]")
    doc.add_paragraph()
    doc.add_paragraph("Technician Signature: _____________________")
    doc.add_paragraph("Customer Sign-off: _____________________")

    doc.save(filepath)


def create_hvac_template(filepath: str):
    doc = Document()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("HVAC PREVENTATIVE MAINTENANCE LOG")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(13, 148, 136) # Teal

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("System Inspection & Performance Verification")
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True

    doc.add_paragraph()

    doc.add_heading("Customer & Unit Specification", level=2)
    doc.add_paragraph("Client Name: _____________________")
    doc.add_paragraph("Facility Address: _____________________")
    doc.add_paragraph("Service Date: [MM/DD/YYYY]")
    doc.add_paragraph("HVAC Unit Model/Serial: _____________________")

    doc.add_paragraph()

    doc.add_heading("System Operating Readings", level=2)
    table = doc.add_table(rows=5, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Operating Metric"
    hdr[1].text = "Specification"
    hdr[2].text = "Reading"

    rdata = [
        ("Refrigerant Suction Pressure (PSI)", "65 - 75 PSI", "_______"),
        ("Refrigerant Discharge Pressure (PSI)", "220 - 260 PSI", "_______"),
        ("Supply Air Temperature (°F)", "52°F - 58°F", "_______"),
        ("Blower Motor Current (Amps)", "< 12.0 A", "_______"),
    ]
    for idx, (m, s, r) in enumerate(rdata, start=1):
        cells = table.rows[idx].cells
        cells[0].text = m
        cells[1].text = s
        cells[2].text = r

    doc.add_paragraph()

    doc.add_heading("Maintenance Tasks Performed", level=2)
    doc.add_paragraph("☐ Air Filters Replaced / Cleaned")
    doc.add_paragraph("☐ Condensate Drain Cleared & Flushed")
    doc.add_paragraph("☐ Fan Belt Tension Adjusted")
    doc.add_paragraph("☐ Thermostat Calibration Verified")

    doc.add_paragraph()

    doc.add_heading("Technician Summary & Approval", level=2)
    doc.add_paragraph("Maintenance Notes: [Summary of work completed and future recommendations]")
    doc.add_paragraph()
    doc.add_paragraph("Lead Technician Signature: _____________________")

    doc.save(filepath)


def create_plumbing_template(filepath: str):
    doc = Document()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("COMMERCIAL PLUMBING SERVICE WORK ORDER")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(37, 99, 235) # Blue

    doc.add_paragraph()

    doc.add_heading("1. Work Order Overview", level=2)
    doc.add_paragraph("Customer Name: _____________________")
    doc.add_paragraph("Site Address: _____________________")
    doc.add_paragraph("Service Date: [MM/DD/YYYY]")
    doc.add_paragraph("Assigned Plumber: _____________________")

    doc.add_paragraph()

    doc.add_heading("2. System Pressure & Flow Inspection", level=2)
    doc.add_paragraph("Static Water Pressure (PSI): _______")
    doc.add_paragraph("Main Shut-off Valve Condition: _____________________")

    doc.add_paragraph()

    doc.add_heading("3. Inspection & Repair Checklist", level=2)
    doc.add_paragraph("☐ Static Water Pressure Verified")
    doc.add_paragraph("☐ Pipe Joint Integrity & Leak Check Confirmed")
    doc.add_paragraph("☐ Drainage Flow Rate Tested")
    doc.add_paragraph("☐ Backflow Preventer Inspection Passed")

    doc.add_paragraph()

    doc.add_heading("4. Work Completed Details", level=2)
    doc.add_paragraph("Summary of Work: [Detailed description of plumbing repairs and materials installed]")
    doc.add_paragraph()
    doc.add_paragraph("Technician Signature: _____________________")
    doc.add_paragraph("Customer Sign-off: _____________________")

    doc.save(filepath)


def create_fire_safety_template(filepath: str):
    doc = Document()

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("FIRE SAFETY & ALARM SYSTEM CERTIFICATION")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(220, 38, 38) # Red

    doc.add_paragraph()

    doc.add_heading("1. Property & Inspection Info", level=2)
    doc.add_paragraph("Property Name: _____________________")
    doc.add_paragraph("Property Address: _____________________")
    doc.add_paragraph("Inspection Date: [MM/DD/YYYY]")
    doc.add_paragraph("Fire Safety Inspector: _____________________")

    doc.add_paragraph()

    doc.add_heading("2. System Readings", level=2)
    doc.add_paragraph("Control Panel Model: _____________________")
    doc.add_paragraph("Standby Battery Voltage (V): _______")

    doc.add_paragraph()

    doc.add_heading("3. Annual Safety Checklist", level=2)
    doc.add_paragraph("☐ Smoke & Heat Detectors Functionality Tested")
    doc.add_paragraph("☐ Manual Pull Stations Operable")
    doc.add_paragraph("☐ Emergency Strobe Lights & Horns Verified")
    doc.add_paragraph("☐ Portable Fire Extinguishers Inspected & Tagged")

    doc.add_paragraph()

    doc.add_heading("4. Certification & Authorization", level=2)
    doc.add_paragraph("Inspector Comments: [Summary of inspection results and compliance status]")
    doc.add_paragraph()
    doc.add_paragraph("Inspector Signature: _____________________")

    doc.save(filepath)


def generate_all_templates():
    out_dir = os.path.join(os.path.dirname(__file__), "..", "uploads", "templates")
    os.makedirs(out_dir, exist_ok=True)

    files = {
        "Electrical Inspection Certificate": os.path.join(out_dir, "Electrical_Inspection_Certificate.docx"),
        "HVAC Maintenance Log": os.path.join(out_dir, "HVAC_Maintenance_Log.docx"),
        "Plumbing Service Work Order": os.path.join(out_dir, "Plumbing_Service_Work_Order.docx"),
        "Fire Safety Certification": os.path.join(out_dir, "Fire_Safety_Certification.docx"),
    }

    create_electrical_template(files["Electrical Inspection Certificate"])
    create_hvac_template(files["HVAC Maintenance Log"])
    create_plumbing_template(files["Plumbing Service Work Order"])
    create_fire_safety_template(files["Fire Safety Certification"])

    print("Generated 4 real .docx templates in uploads/templates/")
    return files

if __name__ == "__main__":
    generate_all_templates()
